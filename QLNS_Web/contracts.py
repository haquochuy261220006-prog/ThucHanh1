"""Soạn hợp đồng lao động bằng AI và xuất ra file Word (.docx).

Nguyên tắc: AI chỉ viết phần ĐIỀU KHOẢN. Các thông tin thật (họ tên, CCCD, lương,
ngày tháng, thông tin công ty...) luôn lấy trực tiếp từ database và được code ghép vào,
nên AI không thể "bịa" sai dữ liệu nhân viên. Dữ liệu cá nhân nhạy cảm (CCCD, SĐT,
địa chỉ, email) không được gửi cho AI.
"""
import io
import re
from datetime import date, datetime
from functools import wraps

from flask import Blueprint, flash, redirect, render_template, request, send_file, session, url_for
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import db
import config

bp = Blueprint("contracts", __name__)

LD_CANCU = ['Căn cứ Bộ luật Lao động số 45/2019/QH14 ngày 20/11/2019;', 'Căn cứ các văn bản hướng dẫn thi hành Bộ luật Lao động hiện hành;', 'Căn cứ nhu cầu sử dụng lao động và thỏa thuận của hai bên.']
DV_CANCU = ["Căn cứ Bộ luật Dân sự số 91/2015/QH13 ngày 24/11/2015;",
            "Căn cứ các văn bản pháp luật có liên quan;",
            "Căn cứ nhu cầu và khả năng của hai bên."]
PARTY_LD = ("NGƯỜI SỬ DỤNG LAO ĐỘNG (Bên A):", "NGƯỜI LAO ĐỘNG (Bên B):", "NGƯỜI LAO ĐỘNG", "NGƯỜI SỬ DỤNG LAO ĐỘNG", "Người lao động")
PARTY_HV = ("NGƯỜI SỬ DỤNG LAO ĐỘNG (Bên A):", "NGƯỜI HỌC VIỆC (Bên B):", "NGƯỜI HỌC VIỆC", "NGƯỜI SỬ DỤNG LAO ĐỘNG", "Người học việc")
PARTY_DV = ("BÊN SỬ DỤNG DỊCH VỤ (Bên A):", "BÊN CUNG CẤP DỊCH VỤ (Bên B):", "BÊN CUNG CẤP DỊCH VỤ", "BÊN SỬ DỤNG DỊCH VỤ", "Bên B")

# Mỗi loại: tiêu đề, mã số HĐ, nhóm (kind), căn cứ pháp lý, cách gọi 2 bên, ghi chú gửi cho AI.
# Muốn thêm loại mới: thêm 1 dòng ở đây + thêm tên vào danh sách 'choice' của hopdong trong app.py.
TYPES = {
    "Thử việc": dict(title="HỢP ĐỒNG THỬ VIỆC", suffix="HĐTV", kind="lao_dong", ten="thử việc", canc=LD_CANCU, party=PARTY_LD,
                     note="Thử việc theo Bộ luật Lao động 2019: thời gian thử việc tối đa tùy trình độ công việc; lương thử việc tối thiểu 85% lương công việc."),
    "Có thời hạn": dict(title="HỢP ĐỒNG LAO ĐỘNG XÁC ĐỊNH THỜI HẠN", suffix="HĐLĐ", kind="lao_dong", ten="lao động", canc=LD_CANCU, party=PARTY_LD,
                        note="Hợp đồng lao động xác định thời hạn, thời hạn tối đa 36 tháng."),
    "Không thời hạn": dict(title="HỢP ĐỒNG LAO ĐỘNG KHÔNG XÁC ĐỊNH THỜI HẠN", suffix="HĐLĐ", kind="lao_dong", ten="lao động", canc=LD_CANCU, party=PARTY_LD,
                           note="Hợp đồng lao động không xác định thời hạn."),
    "Thời vụ": dict(title="HỢP ĐỒNG LAO ĐỘNG THỜI VỤ", suffix="HĐLĐ-TV", kind="lao_dong", ten="lao động thời vụ", canc=LD_CANCU, party=PARTY_LD,
                    note="Công việc theo mùa vụ hoặc một công việc nhất định, thời hạn ngắn (dưới 12 tháng). Theo Bộ luật Lao động 2019 đây là hợp đồng xác định thời hạn; nêu rõ công việc/dự án và điều kiện kết thúc."),
    "Bán thời gian": dict(title="HỢP ĐỒNG LAO ĐỘNG BÁN THỜI GIAN", suffix="HĐLĐ-BTG", kind="lao_dong", ten="lao động bán thời gian", canc=LD_CANCU, party=PARTY_LD,
                          note="Người lao động làm việc không trọn thời gian (thời giờ làm việc ngắn hơn bình thường). Nêu rõ số giờ/ngày, số ngày/tuần, cách tính lương theo giờ hoặc theo tháng tương ứng."),
    "Học việc": dict(title="HỢP ĐỒNG HỌC VIỆC", suffix="HĐHV", kind="hoc_viec", ten="học việc", canc=LD_CANCU, party=PARTY_HV,
                     note="Hợp đồng học nghề/tập nghề tại doanh nghiệp: nêu nghề được học, thời gian, người hướng dẫn, phụ cấp học việc, trách nhiệm đào tạo, không bắt người học làm việc ngoài chương trình, và cam kết xem xét ký hợp đồng lao động khi đạt yêu cầu."),
    "Cộng tác viên": dict(title="HỢP ĐỒNG CỘNG TÁC VIÊN (DỊCH VỤ)", suffix="HĐDV", kind="dich_vu", ten="dịch vụ", canc=DV_CANCU, party=PARTY_DV,
                          note="Đây là hợp đồng dịch vụ/cộng tác theo Bộ luật Dân sự, KHÔNG phải hợp đồng lao động: không có quan hệ lao động, không đóng BHXH, không áp dụng nội quy lao động/nghỉ phép; nêu phạm vi công việc, thù lao, nghiệm thu, thanh toán, thuế, chấm dứt hợp đồng. Gọi hai bên là Bên A (sử dụng dịch vụ) và Bên B (cung cấp dịch vụ)."),
    "Đào tạo": dict(title="HỢP ĐỒNG ĐÀO TẠO", suffix="HĐĐT", kind="dao_tao", ten="đào tạo", canc=LD_CANCU, party=PARTY_LD,
                    note="Hợp đồng đào tạo nâng cao trình độ, kỹ năng nghề giữa doanh nghiệp và người lao động đang làm việc: nội dung khóa đào tạo, thời gian, chi phí đào tạo, quyền lợi trong thời gian đào tạo, cam kết làm việc sau đào tạo và hoàn trả chi phí đào tạo theo quy định của Bộ luật Lao động."),
}
DEFAULT_TYPE = dict(title="HỢP ĐỒNG LAO ĐỘNG", suffix="HĐLĐ", kind="lao_dong", ten="lao động", canc=LD_CANCU, party=PARTY_LD, note="")


def type_info(loai):
    return TYPES.get(loai, DEFAULT_TYPE)


# ----------------------------------------------------------------- tiện ích
def admin_required(f):
    @wraps(f)
    def w(*a, **kw):
        u = session.get("user")
        if not u:
            return redirect(url_for("login"))
        if u.get("role") != "QuanTri":
            flash("Chức năng này dành cho quản trị viên.", "error")
            return redirect(url_for("dashboard"))
        return f(*a, **kw)
    return w


def fmt_date(d):
    if not d:
        return "........"
    if isinstance(d, str):
        d = date.fromisoformat(d[:10])
    return d.strftime("%d/%m/%Y")


def long_date(d):
    d = d or date.today()
    if isinstance(d, str):
        d = date.fromisoformat(d[:10])
    return f"ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"


def money(v):
    return f"{int(v or 0):,}".replace(",", ".")


_DIGITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]


def _read3(n, full):
    tram, chuc, dv = n // 100, (n % 100) // 10, n % 10
    out = []
    if full or tram:
        out.append(_DIGITS[tram] + " trăm")
    if chuc > 1:
        out.append(_DIGITS[chuc] + " mươi")
    elif chuc == 1:
        out.append("mười")
    elif (full or tram) and dv:
        out.append("lẻ")
    if dv:
        if dv == 1 and chuc > 1:
            out.append("mốt")
        elif dv == 5 and chuc > 0:
            out.append("lăm")
        else:
            out.append(_DIGITS[dv])
    return " ".join(out)


def number_to_words(n):
    """Đọc số tiền bằng chữ tiếng Việt, ví dụ 15500000 -> 'Mười lăm triệu năm trăm nghìn đồng'."""
    n = int(n or 0)
    if n == 0:
        return "Không đồng"
    units = ["", " nghìn", " triệu", " tỷ"]
    groups = []
    while n:
        groups.append(n % 1000)
        n //= 1000
    parts = []
    for i in range(len(groups) - 1, -1, -1):
        g = groups[i]
        if g == 0:
            continue
        full = i != len(groups) - 1
        parts.append(_read3(g, full) + (units[i] if i < 4 else ""))
    s = " ".join(parts).strip()
    return s[0].upper() + s[1:] + " đồng"


def months_between(a, b):
    if not a or not b:
        return None
    return (b.year - a.year) * 12 + (b.month - a.month) + (1 if b.day >= a.day else 0)


def load_contract(mahd):
    cols, rows = db.query(
        """SELECT h.MaHD,h.LoaiHD,h.NgayBatDau,h.NgayKetThuc,h.LuongHD,
                  n.MaNV,n.HoTen,n.NgaySinh,n.GioiTinh,n.DiaChi,n.SDT,n.Email,n.CCCD,
                  p.TenPB,c.TenCV
           FROM HopDong h JOIN NhanVien n ON n.MaNV=h.MaNV
           LEFT JOIN PhongBan p ON p.MaPB=n.MaPB LEFT JOIN ChucVu c ON c.MaCV=n.MaCV
           WHERE h.MaHD=?""", (mahd,))
    return dict(zip(cols, rows[0])) if rows else None


# --------------------------------------------------------------------- AI
SYSTEM_PROMPT = """Bạn là chuyên viên pháp chế nhân sự tại Việt Nam, soạn hợp đồng lao động theo pháp luật Việt Nam hiện hành (căn cứ pháp lý chính và loại hợp đồng được nêu trong dữ liệu; hợp đồng lao động theo Bộ luật Lao động 2019 số 45/2019/QH14).

Nhiệm vụ: chỉ soạn PHẦN ĐIỀU KHOẢN của hợp đồng (từ "Điều 1" trở đi). KHÔNG viết quốc hiệu, tiêu đề, thông tin hai bên, phần chữ ký - hệ thống sẽ tự ghép các phần đó.

Quy tắc định dạng (bắt buộc):
- Văn bản thuần, KHÔNG dùng markdown (không **, không #, không bảng, không gạch đầu dòng "-").
- Mỗi điều bắt đầu bằng dòng riêng: "Điều N. Tên điều" (ví dụ "Điều 1. Công việc và địa điểm làm việc").
- Dưới mỗi điều, các khoản đánh số "1.", "2.", ... ; điểm nhỏ đánh "a)", "b)", ...
- Một dòng cho mỗi khoản/điểm. Không dòng trống thừa.
- Điều cuối cùng là "Điều khoản thi hành" (số bản, hiệu lực, giải quyết tranh chấp).

Quy tắc nội dung:
- Chỉ dùng đúng số liệu được cung cấp (chức danh, lương, ngày, thời hạn...). Không tự bịa số tiền, tỷ lệ, số ngày nếu không có trong dữ liệu hoặc quy định pháp luật chắc chắn.
- Gọi hai bên đúng theo dòng "Cách gọi hai bên" trong dữ liệu. Soạn đúng bản chất loại hợp đồng được yêu cầu (lao động, học việc, dịch vụ dân sự, đào tạo); tuyệt đối không biến hợp đồng dịch vụ thành hợp đồng lao động.
- Nếu là hợp đồng thử việc: thời gian thử việc không quá thời hạn tối đa theo luật; tiền lương thử việc ít nhất 85% mức lương của công việc đó; nêu rõ quyền chấm dứt thử việc.
- Phần cần thông tin chưa có thì để chỗ trống "........" thay vì đoán.
- Văn phong hành chính, rõ ràng, trung lập, cân bằng quyền lợi hai bên.
- Chỉ trả về nội dung hợp đồng, không lời dẫn, không giải thích.
- Nếu có công cụ tra cứu web: dùng tối đa vài lần để kiểm tra quy định MỚI NHẤT liên quan (lương tối thiểu vùng, thời gian thử việc tối đa, thời hạn báo trước khi nghỉ việc, bảo hiểm bắt buộc). Chỉ dùng thông tin từ nguồn chính thống (thuvienphapluat.vn, chinhphu.vn, vbpl.vn...). Nếu mức lương trong dữ liệu thấp hơn lương tối thiểu vùng áp dụng, vẫn giữ nguyên số liệu nhưng thêm dòng cuối cùng bắt đầu bằng "LƯU Ý HR:" nêu rõ vấn đề. Mọi trích dẫn số hiệu văn bản phải lấy từ kết quả tra cứu, không tự nhớ."""


def _provider():
    return "openai" if getattr(config, "AI_PROVIDER", "anthropic") == "openai" else "anthropic"


def ai_model():
    return config.OPENAI_MODEL if _provider() == "openai" else config.ANTHROPIC_MODEL


def _ai_available():
    key = config.OPENAI_API_KEY if _provider() == "openai" else config.ANTHROPIC_API_KEY
    return bool(key)


def web_search_supported():
    return _provider() == "anthropic"


def _call_anthropic(system, user_msg, max_tokens, web):
    import anthropic  # nhập khi cần để app vẫn chạy nếu chưa cài thư viện
    client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY, timeout=120.0)
    kw = dict(model=config.ANTHROPIC_MODEL, max_tokens=max_tokens, system=system,
              messages=[{"role": "user", "content": user_msg}])
    if web:
        kw["tools"] = [{"type": "web_search_20250305", "name": "web_search", "max_uses": 4}]
        system += f"\nHôm nay là {date.today().strftime('%d/%m/%Y')}."
        kw["system"] = system
    resp = client.messages.create(**kw)
    blocks = list(resp.content)
    # Khi có tra cứu web, AI có thể nói xen kẽ ("để tôi tìm...") -> chỉ lấy phần chữ SAU lần tra cứu cuối.
    last = max((i for i, b in enumerate(blocks) if "tool_result" in getattr(b, "type", "")), default=-1)
    return "".join(b.text for b in blocks[last + 1:] if getattr(b, "type", "") == "text")


def _call_openai(system, user_msg, max_tokens):
    import json
    import urllib.request
    body = json.dumps({"model": config.OPENAI_MODEL, "max_tokens": max_tokens, "messages": [
        {"role": "system", "content": system}, {"role": "user", "content": user_msg}]}).encode()
    req = urllib.request.Request(config.OPENAI_BASE_URL.rstrip("/") + "/chat/completions", data=body, headers={
        "Content-Type": "application/json", "Authorization": "Bearer " + config.OPENAI_API_KEY})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read().decode("utf-8"))["choices"][0]["message"]["content"]


def _call_ai(system, user_msg, max_tokens=4000, web=False):
    if _provider() == "openai":
        return _call_openai(system, user_msg, max_tokens)
    if web:
        try:
            return _call_anthropic(system, user_msg, max_tokens, True)
        except Exception as e:  # tài khoản chưa bật web search, hoặc lỗi tạm thời -> thử lại không tra cứu
            if "anthropic" in type(e).__module__ or "web" in str(e).lower():
                return _call_anthropic(system, user_msg, max_tokens, False)
            raise
    return _call_anthropic(system, user_msg, max_tokens, False)


def test_connection():
    """Trả về (ok, thông báo)."""
    if not _ai_available():
        return False, "Chưa có API key. Xem mục 5 trong README."
    try:
        out = _call_ai("Chỉ trả lời đúng một từ.", "Trả lời: OK", 20).strip()
        return True, f"Kết nối AI thành công ({_provider()} · {ai_model()}). Phản hồi: {out[:40]}"
    except Exception as e:
        return False, f"Không kết nối được AI: {type(e).__name__}: {e}"


def clean_ai_text(text):
    """Dọn các ký hiệu markdown còn sót, chuẩn hóa dòng trống."""
    text = text.replace("\r\n", "\n")
    text = re.sub(r"```[a-z]*", "", text)
    text = re.sub(r"\*\*|__", "", text)
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*[-•*]\s+", "", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    m = re.search(r"^Điều\s+1\b", text, flags=re.M)  # bỏ lời dẫn trước Điều 1
    if m:
        text = text[m.start():]
    return text.strip()


def describe(c, opts):
    """Mô tả hợp đồng gửi cho AI (không chứa CCCD/SĐT/địa chỉ/email)."""
    months = months_between(c["NgayBatDau"], c["NgayKetThuc"])
    T = type_info(c["LoaiHD"])
    lines = [
        f"Loại hợp đồng: {c['LoaiHD']} ({T['title']})",
        f"Căn cứ pháp lý chính: {'; '.join(T['canc'][:1])}",
        f"Cách gọi hai bên: Bên A = {T['party'][3].capitalize()}, Bên B = {T['party'][4]}",
        f"Lưu ý riêng của loại hợp đồng này: {T['note']}",
        f"Chức danh/vị trí: {c['TenCV'] or '........'}",
        f"Bộ phận/phòng ban: {c['TenPB'] or '........'}",
        f"Ngày bắt đầu: {fmt_date(c['NgayBatDau'])}",
        f"Ngày kết thúc: {fmt_date(c['NgayKetThuc']) if c['NgayKetThuc'] else 'không xác định thời hạn'}"
        + (f" (khoảng {months} tháng)" if months else ""),
        f"Mức lương: {money(c['LuongHD'])} đồng/tháng (bằng chữ: {number_to_words(c['LuongHD'])})",
        f"Địa điểm làm việc: {opts.get('dia_diem') or config.COMPANY.get('dia_chi') or '........'}",
        f"Giờ làm việc: {opts.get('gio_lam') or 'theo nội quy công ty, 8 giờ/ngày'} "
        f"(giờ vào chuẩn {config.GIO_VAO_CHUAN}, {config.CONG_CHUAN} ngày công chuẩn/tháng)",
    ]
    if opts.get("phu_cap"):
        lines.append(f"Phụ cấp/chế độ khác: {opts['phu_cap']}")
    extra = []
    if opts.get("bao_mat"):
        extra.append("điều khoản bảo mật thông tin và bí mật kinh doanh")
    if opts.get("so_huu_tri_tue"):
        extra.append("điều khoản quyền sở hữu trí tuệ đối với sản phẩm công việc")
    if extra:
        lines.append("Bắt buộc có thêm: " + "; ".join(extra))
    if opts.get("yeu_cau"):
        lines.append("Yêu cầu bổ sung của HR: " + opts["yeu_cau"])
    return "\n".join(lines)


def _fb_labor(c, opts):
    thu_viec = c["LoaiHD"] == "Thử việc"
    ten = type_info(c["LoaiHD"])["ten"]
    btg = c["LoaiHD"] == "Bán thời gian"
    if c["NgayKetThuc"]:
        thoi_han = (f"Hợp đồng {ten} có thời hạn từ ngày {fmt_date(c['NgayBatDau'])} "
                    f"đến ngày {fmt_date(c['NgayKetThuc'])}.")
    else:
        thoi_han = (f"Hợp đồng {ten} có hiệu lực từ ngày {fmt_date(c['NgayBatDau'])}"
                    + (" và không xác định thời hạn." if c["LoaiHD"] == "Không thời hạn" else "."))
    dia_diem = opts.get("dia_diem") or config.COMPANY.get("dia_chi") or "........"
    luong_label = "Mức lương thử việc" if thu_viec else "Mức lương"
    out = [
        "Điều 1. Thời hạn và công việc phải làm",
        f"1. {thoi_han}",
        f"2. Chức danh chuyên môn: {c['TenCV'] or '........'}; bộ phận: {c['TenPB'] or '........'}.",
        "3. Công việc phải làm: thực hiện các nhiệm vụ theo bản mô tả công việc và sự phân công của người quản lý trực tiếp.",
        f"4. Địa điểm làm việc: {dia_diem}.",
        "Điều 2. Chế độ làm việc",
        f"1. Thời giờ làm việc: {opts.get('gio_lam') or ('........ giờ/ngày, ........ ngày/tuần (ngắn hơn thời giờ làm việc bình thường)' if btg else '8 giờ/ngày, từ thứ Hai đến thứ Sáu hoặc theo nội quy công ty')}; "
        f"giờ vào làm chuẩn {config.GIO_VAO_CHUAN}.",
        "2. Người lao động được cấp phát công cụ, thiết bị cần thiết để thực hiện công việc và có trách nhiệm bảo quản.",
        "3. Chế độ nghỉ ngơi, nghỉ phép năm, nghỉ lễ, nghỉ việc riêng thực hiện theo Bộ luật Lao động và nội quy công ty.",
        "Điều 3. Tiền lương và chế độ",
        f"1. {luong_label}: {money(c['LuongHD'])} đồng/tháng (bằng chữ: {number_to_words(c['LuongHD'])}).",
        "2. Hình thức trả lương: chuyển khoản vào tài khoản của Người lao động, trả một lần vào tháng sau.",
    ]
    n = 3
    if opts.get("phu_cap"):
        out.append(f"{n}. Phụ cấp và chế độ khác: {opts['phu_cap']}.")
        n += 1
    if not thu_viec:
        out.append(f"{n}. Chế độ bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp thực hiện theo quy định của pháp luật.")
    else:
        out.append(f"{n}. Mức lương thử việc tối thiểu bằng 85% mức lương của công việc làm thử.")
    out += [
        "Điều 4. Nghĩa vụ và quyền lợi của Người lao động",
        "1. Hoàn thành công việc được giao; chấp hành nội quy lao động, quy chế và sự quản lý, điều hành hợp pháp của công ty.",
        "2. Được hưởng lương, chế độ và các quyền lợi khác theo hợp đồng và quy định của pháp luật.",
        "3. Đơn phương chấm dứt hợp đồng theo quy định của pháp luật và báo trước đúng thời hạn quy định.",
    ]
    d = 5
    if opts.get("bao_mat"):
        out += [f"Điều {d}. Bảo mật thông tin",
                "1. Người lao động không được tiết lộ, sao chép hoặc sử dụng thông tin, tài liệu, bí mật kinh doanh của công ty cho mục đích ngoài công việc.",
                "2. Nghĩa vụ bảo mật vẫn có hiệu lực sau khi chấm dứt hợp đồng trong thời hạn theo thỏa thuận hoặc quy định của pháp luật."]
        d += 1
    if opts.get("so_huu_tri_tue"):
        out += [f"Điều {d}. Quyền sở hữu trí tuệ",
                "Sản phẩm, tài liệu, sáng kiến được tạo ra trong quá trình thực hiện công việc theo hợp đồng thuộc quyền sở hữu của công ty, trừ trường hợp pháp luật có quy định khác."]
        d += 1
    out += [
        f"Điều {d}. Nghĩa vụ và quyền hạn của Người sử dụng lao động",
        "1. Bảo đảm việc làm và thực hiện đầy đủ các cam kết trong hợp đồng.",
        "2. Thanh toán đủ, đúng hạn tiền lương và các chế độ cho Người lao động.",
        "3. Có quyền điều hành, giám sát, khen thưởng, xử lý vi phạm theo nội quy lao động và quy định của pháp luật.",
    ]
    d += 1
    if opts.get("yeu_cau"):
        out += [f"Điều {d}. Thỏa thuận khác", opts["yeu_cau"]]
        d += 1
    out += [
        f"Điều {d}. Điều khoản thi hành",
        "1. Hợp đồng này được lập thành 02 bản có giá trị pháp lý như nhau, mỗi bên giữ 01 bản.",
        "2. Những vấn đề chưa ghi trong hợp đồng được thực hiện theo quy định của pháp luật lao động hiện hành.",
        "3. Tranh chấp phát sinh được giải quyết trên cơ sở thương lượng; nếu không thành thì đưa ra cơ quan có thẩm quyền giải quyết.",
    ]
    return "\n".join(out)


def _extras(out, d, opts, subj):
    """Điều khoản tùy chọn (bảo mật, sở hữu trí tuệ, thỏa thuận khác). Trả về số điều tiếp theo."""
    if opts.get("bao_mat"):
        out += [f"Điều {d}. Bảo mật thông tin",
                f"1. {subj} không được tiết lộ, sao chép hoặc sử dụng thông tin, tài liệu, bí mật kinh doanh của Bên còn lại cho mục đích ngoài công việc.",
                "2. Nghĩa vụ bảo mật vẫn có hiệu lực sau khi chấm dứt hợp đồng trong thời hạn theo thỏa thuận hoặc quy định của pháp luật."]
        d += 1
    if opts.get("so_huu_tri_tue"):
        out += [f"Điều {d}. Quyền sở hữu trí tuệ",
                "Sản phẩm, tài liệu, sáng kiến được tạo ra trong quá trình thực hiện hợp đồng thuộc quyền sở hữu của Bên A, trừ trường hợp pháp luật có quy định khác."]
        d += 1
    if opts.get("yeu_cau"):
        out += [f"Điều {d}. Thỏa thuận khác", opts["yeu_cau"]]
        d += 1
    return d


def _closing(out, d):
    out += [f"Điều {d}. Điều khoản thi hành",
            "1. Hợp đồng này được lập thành 02 bản có giá trị pháp lý như nhau, mỗi bên giữ 01 bản.",
            "2. Những vấn đề chưa ghi trong hợp đồng được thực hiện theo quy định của pháp luật hiện hành.",
            "3. Tranh chấp phát sinh được giải quyết trên cơ sở thương lượng; nếu không thành thì đưa ra cơ quan có thẩm quyền giải quyết."]


def _period(c, what):
    if c["NgayKetThuc"]:
        return f"{what} từ ngày {fmt_date(c['NgayBatDau'])} đến ngày {fmt_date(c['NgayKetThuc'])}."
    return f"{what} bắt đầu từ ngày {fmt_date(c['NgayBatDau'])}, chưa xác định ngày kết thúc."


def _fb_apprentice(c, opts):
    dia_diem = opts.get("dia_diem") or config.COMPANY.get("dia_chi") or "........"
    tien = money(c["LuongHD"]) + f" đồng/tháng (bằng chữ: {number_to_words(c['LuongHD'])})" if c["LuongHD"] else "........ đồng/tháng"
    out = [
        "Điều 1. Nội dung và thời gian học việc",
        f"1. Nghề/vị trí học việc: {c['TenCV'] or '........'}; bộ phận: {c['TenPB'] or '........'}.",
        f"2. {_period(c, 'Thời gian học việc')}",
        f"3. Địa điểm học việc: {dia_diem}.",
        "Điều 2. Chương trình và thời gian học",
        f"1. Thời gian học việc: {opts.get('gio_lam') or '8 giờ/ngày hoặc theo lịch của người hướng dẫn'}.",
        "2. Người hướng dẫn: ........ . Bên A xây dựng nội dung, lộ trình học việc và đánh giá kết quả khi kết thúc.",
        "3. Bên A không sử dụng Bên B làm các công việc ngoài chương trình học việc đã thỏa thuận.",
        "Điều 3. Phụ cấp và chế độ",
        f"1. Phụ cấp học việc: {tien}.",
        "2. Trường hợp Bên B trực tiếp hoặc tham gia làm ra sản phẩm hợp pháp thì được trả lương theo mức do hai bên thỏa thuận.",
    ]
    n = 3
    if opts.get("phu_cap"):
        out.append(f"{n}. Phụ cấp và chế độ khác: {opts['phu_cap']}.")
    out += [
        "Điều 4. Nghĩa vụ và quyền lợi của Người học việc",
        "1. Học tập nghiêm túc, chấp hành nội quy, quy định an toàn lao động của Bên A.",
        "2. Được hướng dẫn, cung cấp tài liệu, dụng cụ cần thiết và được bảo đảm an toàn, vệ sinh lao động.",
        "3. Bảo quản tài sản, giữ bí mật thông tin của Bên A trong thời gian học việc.",
        "Điều 5. Nghĩa vụ của Người sử dụng lao động",
        "1. Tổ chức hướng dẫn học việc đúng chương trình; bố trí người hướng dẫn có chuyên môn.",
        "2. Thanh toán đủ, đúng hạn phụ cấp/lương theo hợp đồng.",
        "3. Khi kết thúc học việc, nếu Bên B đạt yêu cầu và Bên A có nhu cầu sử dụng thì hai bên ưu tiên giao kết hợp đồng thử việc hoặc hợp đồng lao động.",
    ]
    d = _extras(out, 6, opts, "Người học việc")
    _closing(out, d)
    return "\n".join(out)


def _fb_service(c, opts):
    dia_diem = opts.get("dia_diem") or "linh hoạt, trực tuyến hoặc tại địa điểm do Bên A đề nghị"
    out = [
        "Điều 1. Nội dung dịch vụ",
        f"1. Bên B thực hiện công việc cộng tác với vai trò: {c['TenCV'] or '........'}; phối hợp với bộ phận: {c['TenPB'] or '........'} của Bên A.",
        "2. Phạm vi công việc cụ thể: ........ (theo yêu cầu bằng văn bản/email của Bên A).",
        f"3. Địa điểm/hình thức thực hiện: {dia_diem}.",
        "Điều 2. Thời hạn hợp đồng",
        f"1. {_period(c, 'Hợp đồng có hiệu lực')}",
        "2. Mỗi bên có quyền chấm dứt hợp đồng bằng cách thông báo trước ít nhất ........ ngày; Bên A thanh toán phần công việc Bên B đã hoàn thành.",
        "Điều 3. Thù lao và thanh toán",
        f"1. Thù lao: {money(c['LuongHD'])} đồng/tháng (bằng chữ: {number_to_words(c['LuongHD'])}), hoặc theo khối lượng công việc được nghiệm thu.",
        "2. Hình thức thanh toán: chuyển khoản, thực hiện trong vòng ........ ngày kể từ khi nghiệm thu kết quả công việc.",
    ]
    n = 3
    if opts.get("phu_cap"):
        out.append(f"{n}. Khoản hỗ trợ khác: {opts['phu_cap']}.")
        n += 1
    out += [
        f"{n}. Thuế: Bên A khấu trừ và kê khai thuế thu nhập cá nhân của Bên B (nếu thuộc diện phải khấu trừ) theo quy định của pháp luật; Bên B chịu trách nhiệm về nghĩa vụ thuế còn lại của mình.",
        "Điều 4. Quyền và nghĩa vụ của Bên B",
        "1. Hoàn thành công việc đúng chất lượng, đúng thời hạn thỏa thuận.",
        "2. Chủ động sắp xếp thời gian và phương thức thực hiện, miễn bảo đảm kết quả công việc.",
        "3. Tự trang bị công cụ, thiết bị trừ trường hợp Bên A có thỏa thuận cung cấp.",
        "4. Hai bên xác nhận đây là hợp đồng dịch vụ theo Bộ luật Dân sự, không phải hợp đồng lao động; Bên A không đóng bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp và không áp dụng chế độ nghỉ phép năm cho Bên B.",
        "Điều 5. Quyền và nghĩa vụ của Bên A",
        "1. Cung cấp đầy đủ thông tin, tài liệu cần thiết để Bên B thực hiện công việc.",
        "2. Nghiệm thu và thanh toán thù lao đầy đủ, đúng hạn.",
        "3. Có quyền yêu cầu Bên B khắc phục khi kết quả không đạt yêu cầu đã thỏa thuận.",
    ]
    d = _extras(out, 6, opts, "Bên B")
    _closing(out, d)
    return "\n".join(out)


def _fb_training(c, opts):
    luong = (f"{money(c['LuongHD'])} đồng/tháng (bằng chữ: {number_to_words(c['LuongHD'])})"
             if c["LuongHD"] else "theo hợp đồng lao động hiện hành")
    out = [
        "Điều 1. Nội dung đào tạo",
        f"1. Khóa đào tạo: ........ ; phục vụ vị trí: {c['TenCV'] or '........'}, bộ phận: {c['TenPB'] or '........'}.",
        "2. Cơ sở đào tạo: ........ ; hình thức đào tạo: ........ .",
        f"3. {_period(c, 'Thời gian đào tạo')}",
        "Điều 2. Chi phí đào tạo",
        "1. Tổng chi phí đào tạo: ........ đồng, bao gồm học phí, tài liệu, chi phí đi lại, lưu trú (nếu có) và các khoản hợp lệ khác có chứng từ.",
        "2. Bên A chi trả chi phí đào tạo nêu trên; chi phí được Bên A xác nhận bằng văn bản sau khi kết thúc khóa học.",
        "Điều 3. Quyền lợi trong thời gian đào tạo",
        f"1. Trong thời gian đào tạo, Bên B được hưởng lương {luong}.",
        "2. Bên B được tạo điều kiện về thời gian để tham gia khóa đào tạo.",
    ]
    if opts.get("phu_cap"):
        out.append(f"3. Hỗ trợ khác: {opts['phu_cap']}.")
    out += [
        "Điều 4. Cam kết làm việc sau đào tạo",
        "1. Sau khi hoàn thành khóa đào tạo, Bên B cam kết tiếp tục làm việc cho Bên A tối thiểu ........ tháng.",
        "2. Bên B áp dụng kiến thức, kỹ năng đã được đào tạo vào công việc theo sự phân công của Bên A.",
        "Điều 5. Hoàn trả chi phí đào tạo",
        "1. Bên B phải hoàn trả chi phí đào tạo nếu tự ý bỏ khóa học, không hoàn thành khóa học do lỗi của mình, hoặc đơn phương chấm dứt hợp đồng lao động trái cam kết, hoặc bị xử lý kỷ luật sa thải trong thời gian cam kết.",
        "2. Mức hoàn trả được tính theo tỷ lệ thời gian còn lại chưa thực hiện cam kết và theo quy định của Bộ luật Lao động, nhưng không vượt quá chi phí đào tạo thực tế đã được xác nhận.",
        "Điều 6. Nghĩa vụ của Người lao động",
        "1. Tham gia đầy đủ, nghiêm túc và nỗ lực hoàn thành khóa đào tạo; báo cáo kết quả học tập cho Bên A.",
        "2. Bảo quản tài liệu, thông tin được cung cấp và không sử dụng cho mục đích ngoài công việc.",
        "Điều 7. Nghĩa vụ của Người sử dụng lao động",
        "1. Thanh toán chi phí đào tạo và lương, chế độ trong thời gian đào tạo đúng thỏa thuận.",
        "2. Bố trí công việc phù hợp với trình độ, kỹ năng Bên B đạt được sau đào tạo.",
    ]
    d = _extras(out, 8, opts, "Người lao động")
    _closing(out, d)
    return "\n".join(out)


def fallback_text(c, opts):
    """Điều khoản chuẩn dùng khi chưa cấu hình AI hoặc AI lỗi, theo từng loại hợp đồng."""
    return {"hoc_viec": _fb_apprentice, "dich_vu": _fb_service, "dao_tao": _fb_training}.get(
        type_info(c["LoaiHD"])["kind"], _fb_labor)(c, opts)


def generate_clauses(c, opts, web=False):
    """Trả về (nội dung, thông báo). Tự động quay về mẫu chuẩn nếu AI không dùng được."""
    if not _ai_available():
        return fallback_text(c, opts), ("Chưa cấu hình API key nên đang dùng mẫu điều khoản chuẩn "
                                        "(chưa có AI). Xem hướng dẫn trong README.")
    try:
        text = clean_ai_text(_call_ai(SYSTEM_PROMPT, "Hãy soạn điều khoản hợp đồng với thông tin sau:\n\n" + describe(c, opts), web=web))
        if not re.search(r"^Điều\s+1\b", text, flags=re.M):
            raise ValueError("AI trả về nội dung không đúng định dạng")
        return text, None
    except Exception as e:
        return fallback_text(c, opts), f"Không gọi được AI ({type(e).__name__}: {e}). Đã dùng mẫu điều khoản chuẩn."


def revise_clauses(c, opts, current, instruction, web=False):
    if not _ai_available():
        raise RuntimeError("Chưa cấu hình API key.")
    msg = (f"Thông tin hợp đồng:\n{describe(c, opts)}\n\nNội dung điều khoản hiện tại:\n{current}\n\n"
           f"Yêu cầu chỉnh sửa: {instruction}\n\nHãy trả về TOÀN BỘ phần điều khoản sau khi chỉnh sửa, "
           f"giữ nguyên định dạng quy định.")
    text = clean_ai_text(_call_ai(SYSTEM_PROMPT, msg, web=web))
    if not re.search(r"^Điều\s+1\b", text, flags=re.M):
        raise ValueError("AI trả về nội dung không đúng định dạng")
    return text


# ------------------------------------------------------------------- Word
def _font(run, size=13, bold=False, italic=False, underline=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold, run.italic, run.underline = bold, italic, underline


def _para(doc_or_cell, text="", align=None, bold=False, italic=False, size=13,
          before=0, after=4, indent=None, first=None, underline=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.2
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if text:
        _font(p.add_run(text), size, bold, italic, underline)
    return p


def _rich(p, parts, size=13):
    """parts: list[(text, bold)]"""
    for t, b in parts:
        _font(p.add_run(t), size, b)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _first_cell_para(cell):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    return p


def build_docx(c, so_hd, ngay_ky, clauses):
    co = config.COMPANY
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # Quốc hiệu
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(t)
    left, right = t.rows[0].cells
    left.width, right.width = Cm(7), Cm(9)
    p = _first_cell_para(left)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run(co.get("ten", "CÔNG TY ........").upper()), 12, True)
    p = _first_cell_para(right)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"), 12, True)
    for txt, b, u in (("Độc lập – Tự do – Hạnh phúc", True, False), ("────────────", False, False)):
        q = right.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        _font(q.add_run(txt), 13, b)

    loai = c["LoaiHD"]
    T = type_info(loai)
    _para(doc, T["title"], WD_ALIGN_PARAGRAPH.CENTER, True, size=15, before=14, after=2)
    _para(doc, f"Số: {so_hd}", WD_ALIGN_PARAGRAPH.CENTER, italic=True, after=10)

    for line in T["canc"]:
        _para(doc, line, italic=True, after=2, first=0.8)
    _para(doc, f"Hôm nay, {long_date(ngay_ky)}, tại {co.get('dia_chi') or '........'}, chúng tôi gồm:",
          before=6, after=6, first=0.8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # Bên A
    _para(doc, T["party"][0], bold=True, before=4, after=3)
    for label, val in (("Tên công ty", co.get("ten", "")), ("Địa chỉ", co.get("dia_chi", "")),
                       ("Mã số thuế", co.get("mst", "")), ("Điện thoại", co.get("dien_thoai", "")),
                       ("Đại diện", co.get("dai_dien", "")), ("Chức vụ", co.get("chuc_vu", ""))):
        _rich(_para(doc, after=2, indent=0.8), [(f"{label}: ", False), (val or "........", label == "Đại diện")])

    # Bên B
    gt = c["GioiTinh"] or ""
    ong_ba = "Bà" if gt == "Nữ" else ("Ông" if gt == "Nam" else "Ông/Bà")
    _para(doc, T["party"][1], bold=True, before=8, after=3)
    _rich(_para(doc, after=2, indent=0.8), [(f"{ong_ba}: ", False), ((c["HoTen"] or "........").upper(), True)])
    for label, val in (("Ngày sinh", fmt_date(c["NgaySinh"]) if c["NgaySinh"] else None),
                       ("Số CCCD", c["CCCD"]), ("Địa chỉ liên hệ", c["DiaChi"]),
                       ("Điện thoại", c["SDT"]), ("Email", c["Email"])):
        _rich(_para(doc, after=2, indent=0.8), [(f"{label}: ", False), (val or "........", False)])
    _para(doc, "Hai bên cùng thỏa thuận ký kết hợp đồng và cam kết thực hiện đúng các điều khoản sau đây:",
          before=8, after=6, first=0.8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # Điều khoản
    for raw in clauses.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if re.match(r"^Điều\s+\d+", line):
            _para(doc, line, bold=True, before=8, after=3).paragraph_format.keep_with_next = True
        elif re.match(r"^[a-zđ]\)\s", line):
            _para(doc, line, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.4, after=3)
        else:
            _para(doc, line, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.8, after=3)

    # Chữ ký
    _para(doc, f"{co.get('dia_diem_ky') or ''}{', ' if co.get('dia_diem_ky') else ''}{long_date(ngay_ky)}",
          WD_ALIGN_PARAGRAPH.RIGHT, italic=True, before=12, after=6).paragraph_format.keep_with_next = True
    sig = doc.add_table(rows=1, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(sig)
    for cell, title, note, name in (
        (sig.rows[0].cells[0], T["party"][2], "(Ký, ghi rõ họ tên)", c["HoTen"]),
        (sig.rows[0].cells[1], T["party"][3], "(Ký, đóng dấu, ghi rõ họ tên)", co.get("dai_dien", "")),
    ):
        p = _first_cell_para(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(title), 13, True)
        q = cell.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.keep_with_next = True
        _font(q.add_run(note), 12, italic=True)
        for _ in range(4):
            cell.add_paragraph().paragraph_format.keep_with_next = True
        r = cell.add_paragraph()
        r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(r.add_run(name or ""), 13, True)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ------------------------------------------------------------------ routes
def _opts_from_form(f):
    return {
        "dia_diem": f.get("dia_diem", "").strip(),
        "gio_lam": f.get("gio_lam", "").strip(),
        "phu_cap": f.get("phu_cap", "").strip(),
        "yeu_cau": f.get("yeu_cau", "").strip(),
        "bao_mat": bool(f.get("bao_mat")),
        "so_huu_tri_tue": bool(f.get("so_huu_tri_tue")),
        "web": bool(f.get("web")) and web_search_supported(),
    }


def _default_so_hd(c):
    return f"{(c['MaHD'] or c['MaNV']):03d}/{date.today().year}/{type_info(c['LoaiHD'])['suffix']}"


def _contract_list():
    return db.query(
        "SELECT h.MaHD,n.HoTen,h.LoaiHD,h.NgayBatDau FROM HopDong h JOIN NhanVien n ON n.MaNV=h.MaNV ORDER BY h.MaHD DESC")[1]


def _employee_list():
    return db.query(
        """SELECT n.MaNV,n.HoTen,ISNULL(c.TenCV,N''),n.LuongCoBan FROM NhanVien n
           LEFT JOIN ChucVu c ON c.MaCV=n.MaCV WHERE n.TrangThai=N'Đang làm' ORDER BY n.HoTen""")[1]


def _new_form(f=None):
    f = f or {}
    return {"manv": f.get("manv", ""), "loai": f.get("loai") or "Có thời hạn",
            "ngay_bd": f.get("ngay_bd") or date.today().isoformat(),
            "ngay_kt": f.get("ngay_kt", ""), "luong": f.get("luong", "")}


def load_new_contract(f):
    """Dựng thông tin hợp đồng từ nhân viên + loại HĐ chọn trên form (chưa lưu vào database)."""
    manv = f.get("manv", type=int)
    loai = f.get("loai")
    if not manv or loai not in TYPES:
        return None, "Vui lòng chọn nhân viên và loại hợp đồng."
    cols, rows = db.query(
        """SELECT n.MaNV,n.HoTen,n.NgaySinh,n.GioiTinh,n.DiaChi,n.SDT,n.Email,n.CCCD,p.TenPB,c.TenCV
           FROM NhanVien n LEFT JOIN PhongBan p ON p.MaPB=n.MaPB LEFT JOIN ChucVu c ON c.MaCV=n.MaCV
           WHERE n.MaNV=?""", (manv,))
    if not rows:
        return None, "Không tìm thấy nhân viên."
    try:
        bd = date.fromisoformat(f.get("ngay_bd") or date.today().isoformat())
        kt = date.fromisoformat(f["ngay_kt"]) if f.get("ngay_kt") and loai != "Không thời hạn" else None
        luong = int(float(f.get("luong") or 0))
    except ValueError:
        return None, "Ngày hoặc lương không hợp lệ."
    if kt and kt < bd:
        return None, "Ngày kết thúc phải sau ngày bắt đầu."
    c = dict(zip(cols, rows[0]))
    c.update(MaHD=None, LoaiHD=loai, NgayBatDau=bd, NgayKetThuc=kt, LuongHD=luong)
    return c, None


@bp.route("/hopdong/soan", methods=["GET", "POST"])
@bp.route("/hopdong/soan/<int:mahd>", methods=["GET", "POST"])
@admin_required
def soan(mahd=None):
    ctx = dict(ai_on=_ai_available(), model=ai_model(), web_ok=web_search_supported(), contracts=[], employees=[],
               types=list(TYPES), mode="co", form=_new_form(), c=None, text=None,
               opts={"web": bool(getattr(config, "AI_WEB_SEARCH", False))}, so_hd="",
               ngay_ky=date.today().isoformat(), company_ok=bool(config.COMPANY.get("ten")), admin=True)
    try:
        ctx["contracts"] = _contract_list()
        ctx["employees"] = _employee_list()
    except Exception as e:
        flash("Lỗi truy vấn dữ liệu: " + str(e), "error")
        return render_template("contract_ai.html", **ctx)

    if request.method == "GET":
        mahd = mahd or request.args.get("mahd", type=int)
        if mahd:
            c = load_contract(mahd)
            if c:
                ctx.update(c=c, so_hd=_default_so_hd(c))
            else:
                flash("Không tìm thấy hợp đồng.", "error")
        else:
            ctx["mode"] = "moi"  # mặc định soạn mới; chọn "Hợp đồng đã có" khi cần
            if request.args.get("manv"):
                ctx["form"] = _new_form({"manv": request.args.get("manv")})
        return render_template("contract_ai.html", **ctx)

    f = request.form
    mode = "moi" if f.get("mode") == "moi" else "co"
    ctx["mode"] = mode
    if mode == "moi":
        ctx["form"] = _new_form(f)
        c, err = load_new_contract(f)
    else:
        mahd = f.get("MaHD", type=int) or mahd
        c = load_contract(mahd) if mahd else None
        err = None if c else "Vui lòng chọn hợp đồng cần soạn."
    if err:
        flash(err, "error")
        return render_template("contract_ai.html", **ctx)

    opts = _opts_from_form(f)
    so_hd = f.get("so_hd", "").strip() or _default_so_hd(c)
    ngay_ky = f.get("ngay_ky") or date.today().isoformat()
    ctx.update(c=c, opts=opts, so_hd=so_hd, ngay_ky=ngay_ky)
    action = f.get("action")
    current = f.get("noi_dung", "")

    if action == "generate":
        text, warn = generate_clauses(c, opts, opts["web"])
        ctx["text"] = text
        flash(warn, "error") if warn else flash("AI đã soạn xong bản nháp. Hãy đọc lại và chỉnh sửa trước khi xuất Word.", "success")
    elif action == "revise":
        ctx["text"] = current
        instruction = f.get("chinh_sua", "").strip()
        if not instruction:
            flash("Hãy nhập yêu cầu chỉnh sửa cho AI.", "error")
        else:
            try:
                ctx["text"] = revise_clauses(c, opts, current, instruction, opts["web"])
                flash("AI đã chỉnh sửa theo yêu cầu.", "success")
            except Exception as e:
                flash(f"Không chỉnh sửa được bằng AI: {e}", "error")
    elif action == "download":
        ctx["text"] = current
        if not current.strip():
            flash("Chưa có nội dung điều khoản để xuất.", "error")
            return render_template("contract_ai.html", **ctx)
        if mode == "moi" and f.get("luu"):
            try:
                db.execute("INSERT INTO HopDong(MaNV,LoaiHD,NgayBatDau,NgayKetThuc,LuongHD) VALUES(?,?,?,?,?)",
                           (c["MaNV"], c["LoaiHD"], c["NgayBatDau"], c["NgayKetThuc"], c["LuongHD"]))
                flash("Đã lưu hợp đồng vào danh sách Hợp đồng.", "success")
            except Exception as e:
                flash("Không lưu được hợp đồng vào hệ thống: " + str(e), "error")
        buf = build_docx(c, so_hd, ngay_ky, current)
        safe = re.sub(r"[^\w\-]+", "_", c["HoTen"] or "NV", flags=re.U)
        return send_file(buf, as_attachment=True, download_name=f"HopDong_{c['MaHD'] or 'moi'}_{safe}.docx",
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    return render_template("contract_ai.html", **ctx)


@bp.route("/hopdong/ai-test")
@admin_required
def ai_test():
    ok, msg = test_connection()
    flash(msg, "success" if ok else "error")
    return redirect(url_for("contracts.soan"))
